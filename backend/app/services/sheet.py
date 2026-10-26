"""Импорт презентаций из Google-таблицы: 1-й столбец — название, 2-й — дата выступления."""
import asyncio
import csv
import hashlib
import io
import logging
import re
import urllib.request
from datetime import datetime, timedelta, timezone

from sqlalchemy import select
from sqlalchemy.orm import Session

from app.config import get_settings
from app.database import SessionLocal
from app.models import SheetRow, Task
from app.services import tasks as svc

log = logging.getLogger("sheet")

DATE_RE = re.compile(r"(\d{1,4})[./-](\d{1,2})(?:[./-](\d{1,4}))?(?:[ T,]+(\d{1,2})[:.](\d{2}))?")
SHEET_ID_RE = re.compile(r"/spreadsheets/d/([\w-]+)")
GID_RE = re.compile(r"[#&?]gid=(\d+)")


def csv_url(url: str) -> str:
    """Ссылка на таблицу (…/edit#gid=0) -> ссылка на CSV-экспорт; прямую ссылку на экспорт возвращаем как есть."""
    m = SHEET_ID_RE.search(url)
    if not m or "/export" in url or "/pub" in url:
        return url
    gid = GID_RE.search(url)
    return f"https://docs.google.com/spreadsheets/d/{m.group(1)}/export?format=csv&gid={gid.group(1) if gid else 0}"


def parse_date(raw: str) -> datetime | None:
    """dd.mm.yyyy / dd/mm/yyyy / yyyy-mm-dd / dd.mm (год по умолчанию из настроек), опционально с временем.
    Только дата -> 12:00 UTC (не уезжает на соседний день)."""
    m = DATE_RE.search(raw)
    if not m:
        return None
    a, b = int(m.group(1)), int(m.group(2))
    if m.group(3) is None:
        y, mo, d = get_settings().sheet_default_year, b, a
    else:
        c = int(m.group(3))
        y, mo, d = (a, b, c) if len(m.group(1)) == 4 else (c, b, a)
    if y < 100:
        y += 2000
    hh, mm = (int(m.group(4)), int(m.group(5))) if m.group(4) else (12, 0)
    try:
        return datetime(y, mo, d, hh, mm, tzinfo=timezone.utc)
    except ValueError:
        return None


def parse_rows(text: str, links: dict[int, str] | None = None) -> list[tuple[str, str, datetime | None, str]]:
    rows = list(csv.reader(io.StringIO(text)))
    out = []
    for i, row in enumerate(rows):
        title = row[0].strip() if row else ""
        raw = row[1].strip() if len(row) > 1 else ""
        link = row[2].strip() if len(row) > 2 else ""
        link = (links or {}).get(i) or link  # CSV отдаёт только видимый текст ячейки, а настоящий URL лежит в гиперссылке
        due = parse_date(raw)
        if not title or (i == 0 and due is None):  # пустая строка или шапка таблицы
            continue
        out.append((title[:200], raw, due, link))
    return out


def row_key(title: str, raw_date: str) -> str:
    return hashlib.sha1(f"{title.casefold()}|{raw_date}".encode()).hexdigest()


LINK_PREFIX = "Презентация: "


def _link_line(link: str) -> str:
    return f"{LINK_PREFIX}{link}"


def import_rows(db: Session, rows: list[tuple[str, str, datetime | None, str]]) -> int:
    settings = get_settings()
    board = svc.get_board_by_name(db, settings.sheet_board)
    column = svc.get_or_create_column(db, settings.sheet_column, board.id)
    known = {k for (k,) in db.query(SheetRow.key)}
    created = 0
    for title, raw, due, link in rows:
        key = row_key(title, raw)
        if key in known:
            # дозаполняем то, что при первом импорте не распозналось или отсутствовало (дата без года, ссылка в 3-м столбце)
            row = db.get(SheetRow, key)
            task = row and row.task_id and db.get(Task, row.task_id)
            if task and task.source == "sheet":
                if due and task.due_at is None:
                    task.due_at = due
                    task.description = task.description.replace(f"Дата выступления: {raw}", "").strip()
                if link and _link_line(link) not in task.description.splitlines():
                    kept = [ln for ln in task.description.splitlines() if not ln.startswith(LINK_PREFIX)]
                    task.description = "\n".join([*kept, _link_line(link)]).strip()
            continue
        known.add(key)
        task = Task(
            title=title,
            description="\n".join(
                filter(None, [None if due or not raw else f"Дата выступления: {raw}", _link_line(link) if link else None])
            ),
            column_id=column.id,
            position=svc.next_position(db, column.id),
            due_at=due,
            source="sheet",
            author_name="Google Sheets",
        )
        db.add(task)
        db.flush()
        db.add(SheetRow(key=key, task_id=task.id))
        created += 1
    db.commit()
    return created


def archive_past(db: Session) -> int:
    """Задачи из таблицы, которые ещё лежат в To Do и у которых дата прошла, переносит в Done.
    Задачи, которые пользователь уже сам переложил в другую колонку, не трогаем."""
    settings = get_settings()
    board = svc.get_board_by_name(db, settings.sheet_board)
    todo = svc.get_or_create_column(db, settings.sheet_column, board.id)
    done = svc.get_or_create_column(db, settings.sheet_done_column, board.id)
    past = db.scalars(
        select(Task)
        .where(Task.column_id == todo.id, Task.source == "sheet", Task.due_at.is_not(None), Task.due_at < datetime.now(timezone.utc))
        .order_by(Task.due_at)
    ).all()
    for t in past:
        t.column_id = done.id
        t.position = svc.next_position(db, done.id)
        db.flush()
    if past:
        svc._renumber(db, todo.id)
    db.commit()
    return len(past)


def _fetch(url: str) -> str:
    req = urllib.request.Request(url, headers={"User-Agent": "taskgram"})
    with urllib.request.urlopen(req, timeout=20) as resp:
        return resp.read().decode("utf-8-sig")


def parse_xlsx_links(data: bytes) -> dict[int, str]:
    """Гиперссылки из 3-го столбца: {индекс строки с нуля: URL}."""
    from openpyxl import load_workbook

    ws = load_workbook(io.BytesIO(data), read_only=False).active
    out = {}
    for i, row in enumerate(ws.iter_rows(min_col=3, max_col=3)):
        cell = row[0]
        target = cell.hyperlink.target if cell.hyperlink and cell.hyperlink.target else None
        if target:
            out[i] = target
    return out


def fetch_links(url: str) -> dict[int, str]:
    try:
        req = urllib.request.Request(url.replace("format=csv", "format=xlsx"), headers={"User-Agent": "taskgram"})
        with urllib.request.urlopen(req, timeout=30) as resp:
            return parse_xlsx_links(resp.read())
    except Exception as e:  # noqa: BLE001 — без гиперссылок остаётся текст ячейки
        log.warning("sheet: не удалось прочитать гиперссылки: %s", e)
        return {}


def sync_once() -> tuple[int, int]:
    """Возвращает (создано задач, перенесено в Done)."""
    url = csv_url(get_settings().sheet_url)
    text = _fetch(url)
    if text.lstrip().lower().startswith(("<!doctype", "<html")):
        raise RuntimeError("таблица недоступна: откройте доступ «у кого есть ссылка — читатель»")
    with SessionLocal() as db:
        created = import_rows(db, parse_rows(text, fetch_links(url)))
        return created, archive_past(db)


async def sync_loop() -> None:
    settings = get_settings()
    while True:
        try:
            created, archived = await asyncio.to_thread(sync_once)
            if created or archived:
                log.info("sheet: создано задач: %d, в Done: %d", created, archived)
        except Exception as e:  # noqa: BLE001 — сеть/таблица не должны ронять сервер
            log.warning("sheet: %s", e)
        await asyncio.sleep(max(settings.sheet_poll_seconds, 10))
